"""
KNAVE — Descargas y conversiones de archivos.
Backend FastAPI. Diseñado para fallar con gracia: cada error llega al
frontend como un mensaje legible, nunca como un cuelgue silencioso.

Uso:  python main.py   (o: uvicorn main:app --host 127.0.0.1 --port 8666)
Modo web multiusuario: KNAVE_SERVER=1 (opcional KNAVE_PASSWORD) — ver DESPLIEGUE.md.
Requiere: ffmpeg (audio/video). LibreOffice para conversiones a PDF.
"""

from __future__ import annotations

import base64
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import unicodedata
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urlparse

from fastapi import (FastAPI, File, Form, HTTPException, Request, Response,
                     UploadFile)
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

# ----------------------------------------------------------------------------
# Rutas base — compatibles con la app empaquetada (PyInstaller)
# ----------------------------------------------------------------------------
if getattr(sys, "frozen", False):
    BUNDLE_DIR = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
else:
    BUNDLE_DIR = Path(__file__).parent.resolve()

STATIC_DIR = BUNDLE_DIR / "static"

CONFIG_DIR = Path.home() / ".knave"          # datos de usuario, siempre escribible
CONFIG_DIR.mkdir(exist_ok=True)
CONFIG_FILE = CONFIG_DIR / "config.json"
WORKSPACE = CONFIG_DIR / "workspace"         # borrador temporal (se limpia solo)
WORKSPACE.mkdir(exist_ok=True)

# ----------------------------------------------------------------------------
# Modo servidor (despliegue web multiusuario) — se activa con KNAVE_SERVER=1
# Cada visitante recibe una sesión anónima por cookie: biblioteca, trabajos y
# ajustes quedan aislados; nadie puede ver lo de otro.
# ----------------------------------------------------------------------------
SERVER_MODE = os.environ.get("KNAVE_SERVER") == "1"
SITE_PASSWORD = os.environ.get("KNAVE_PASSWORD", "")
SESSION_TTL_HOURS = int(os.environ.get("KNAVE_TTL_HOURS", "24"))
SESSIONS_DIR = CONFIG_DIR / "sessions"
if SERVER_MODE:
    SESSIONS_DIR.mkdir(exist_ok=True)


def _session_dir(sid: str) -> Path:
    d = SESSIONS_DIR / sid
    d.mkdir(parents=True, exist_ok=True)
    return d

# ----------------------------------------------------------------------------
# Constantes
# ----------------------------------------------------------------------------
MAX_UPLOAD_MB = int(os.environ.get("KNAVE_MAX_MB", "500"))
SCRATCH_TTL_SECONDS = 6 * 60 * 60            # los borradores caducan a las 6 h
DOWNLOAD_TIMEOUT = 60 * 60
SOFFICE_TIMEOUT = 300
MAX_PDF_PAGES = 500                          # tope al rasterizar/dividir PDFs
MAX_WORKERS = 3

AUDIO_QUALITIES = {"128", "192", "320"}
AUDIO_FORMATS = {"mp3", "m4a", "wav", "flac", "opus"}
VIDEO_QUALITIES = {"480", "720", "1080", "1440", "2160", "best"}

IMG_EXT = {"jpg", "jpeg", "png", "webp", "bmp", "tiff", "tif", "gif"}
OFFICE_EXT = {"doc", "docx", "odt", "rtf", "txt", "ppt", "pptx", "odp",
              "xls", "xlsx", "csv"}

# Operaciones del conversor: qué aceptan y cuántos archivos admiten.
OPERATIONS: dict[str, dict] = {
    "office2pdf":  {"label": "A PDF",           "exts": OFFICE_EXT, "min": 1, "max": 1},
    "pdf2docx":    {"label": "A Word",          "exts": {"pdf"},    "min": 1, "max": 1},
    "pdf2pptx":    {"label": "A PowerPoint",    "exts": {"pdf"},    "min": 1, "max": 1},
    "pdf2img":     {"label": "A imágenes",      "exts": {"pdf"},    "min": 1, "max": 1},
    "pdfsplit":    {"label": "Separar páginas", "exts": {"pdf"},    "min": 1, "max": 1},
    "pdfcompress": {"label": "Comprimir",       "exts": {"pdf"},    "min": 1, "max": 1},
    "img2pdf":     {"label": "Unir en PDF",     "exts": IMG_EXT,    "min": 1, "max": 200},
    "pdfmerge":    {"label": "Combinar PDFs",   "exts": {"pdf"},    "min": 2, "max": 200},
}

app = FastAPI(title="KNAVE", docs_url=None, redoc_url=None)


@app.middleware("http")
async def _access_mw(request: Request, call_next):
    if SITE_PASSWORD and request.url.path != "/healthz":
        header = request.headers.get("authorization", "")
        granted = False
        if header.startswith("Basic "):
            try:
                granted = (base64.b64decode(header[6:]).decode("utf-8")
                           .split(":", 1)[1] == SITE_PASSWORD)
            except Exception:  # noqa: BLE001
                granted = False
        if not granted:
            return Response("KNAVE: se necesita la contraseña.", status_code=401,
                            headers={"WWW-Authenticate": 'Basic realm="KNAVE"'})
    sid, fresh = request.cookies.get("knave_sid", ""), False
    if SERVER_MODE and not re.fullmatch(r"[0-9a-f]{32}", sid or ""):
        sid, fresh = uuid.uuid4().hex, True
    request.state.sid = sid if SERVER_MODE else "local"
    response = await call_next(request)
    if SERVER_MODE and fresh:
        response.set_cookie("knave_sid", sid, max_age=180 * 24 * 3600,
                            httponly=True, samesite="lax")
    return response

# ----------------------------------------------------------------------------
# Ajustes persistentes
# ----------------------------------------------------------------------------
def _default_output() -> Path:
    downloads = Path.home() / "Downloads"
    return (downloads if downloads.exists() else Path.home()) / "KNAVE"


DEFAULT_SETTINGS = {
    "output_dir": str(_default_output()),
    "embed_thumbnail": True,
    "download_subs": False,
    "auto_open_folder": False,
    "audio_format": "mp3",
}


def _settings_file(sid: str) -> Path:
    if SERVER_MODE and sid != "local":
        return _session_dir(sid) / "config.json"
    return CONFIG_FILE


def load_settings(sid: str = "local") -> dict:
    data = dict(DEFAULT_SETTINGS)
    try:
        data.update(json.loads(_settings_file(sid).read_text("utf-8")))
    except Exception:  # noqa: BLE001 — config ausente o corrupta → por defecto
        pass
    return data


def save_settings(data: dict, sid: str = "local") -> dict:
    current = load_settings(sid)
    current.update({k: v for k, v in data.items() if k in DEFAULT_SETTINGS})
    _settings_file(sid).write_text(json.dumps(current, indent=2), "utf-8")
    return current


def output_dir(sid: str = "local") -> Path:
    if SERVER_MODE and sid != "local":
        d = _session_dir(sid) / "library"   # biblioteca privada de la sesión
    else:
        d = Path(load_settings(sid)["output_dir"]).expanduser()
    d.mkdir(parents=True, exist_ok=True)
    return d


def _cookies_path(sid: str = "local") -> Path:
    """cookies.txt del usuario (opcional): rescata descargas bloqueadas."""
    if SERVER_MODE and sid != "local":
        return _session_dir(sid) / "cookies.txt"
    return CONFIG_DIR / "cookies.txt"


# ----------------------------------------------------------------------------
# Registro de trabajos + cola de notificaciones (la consume la bandeja)
# ----------------------------------------------------------------------------
JOBS: dict[str, dict] = {}
JOBS_LOCK = threading.Lock()
POOL = ThreadPoolExecutor(max_workers=MAX_WORKERS)

_NOTIFICATIONS: list[dict] = []
_NOTIF_LOCK = threading.Lock()


def _notify(title: str, message: str) -> None:
    if SERVER_MODE:          # los avisos de bandeja no aplican en la web
        return
    with _NOTIF_LOCK:
        _NOTIFICATIONS.append({"title": title, "message": message})


def pop_notifications() -> list[dict]:
    """desktop.py lo llama en bucle para mostrar avisos nativos."""
    with _NOTIF_LOCK:
        items = list(_NOTIFICATIONS)
        _NOTIFICATIONS.clear()
    return items


class Cancelled(Exception):
    """Se lanza cuando el usuario cancela un trabajo en curso."""


def _new_job(kind: str, label: str, sid: str = "local") -> str:
    jid = uuid.uuid4().hex[:12]
    with JOBS_LOCK:
        JOBS[jid] = {
            "id": jid, "kind": kind, "label": label, "sid": sid,
            "status": "queued", "progress": 0, "detail": "En cola…",
            "filename": None, "filepath": None, "error": None,
            "cancelled": False, "proc": None, "created": time.time(),
        }
    return jid


def _update(jid: str, **fields) -> None:
    with JOBS_LOCK:
        job = JOBS.get(jid)
        if job is not None:
            job.update(fields)


def _is_cancelled(jid: str) -> bool:
    with JOBS_LOCK:
        job = JOBS.get(jid)
        return bool(job and job["cancelled"])


def _check_cancel(jid: str) -> None:
    if _is_cancelled(jid):
        raise Cancelled()


def _fail(jid: str, message: str) -> None:
    _update(jid, status="error", error=message, detail=message, progress=100)


def _job_dir(jid: str) -> Path:
    d = WORKSPACE / jid
    d.mkdir(parents=True, exist_ok=True)
    return d


def _sanitize_name(name: str, max_len: int = 150) -> str:
    name = unicodedata.normalize("NFKC", name)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    return (name or "archivo")[:max_len]


def _unique_in(folder: Path, name: str) -> Path:
    """Evita pisar archivos ya existentes en la carpeta de salida."""
    target = folder / name
    if not target.exists():
        return target
    stem, suffix = target.stem, target.suffix
    i = 2
    while (folder / f"{stem} ({i}){suffix}").exists():
        i += 1
    return folder / f"{stem} ({i}){suffix}"


def _publish(jid: str, produced: list[Path], zip_stem: str) -> None:
    """Mueve el resultado a la carpeta permanente del usuario y cierra el job."""
    files = [p for p in produced if p.is_file()]
    if not files:
        _fail(jid, "El proceso terminó pero no produjo ningún archivo.")
        return
    with JOBS_LOCK:
        sid = JOBS.get(jid, {}).get("sid", "local")
    out = output_dir(sid)
    if len(files) == 1:
        final = _unique_in(out, files[0].name)
        shutil.move(str(files[0]), final)
    else:
        final = _unique_in(out, f"{_sanitize_name(zip_stem)}.zip")
        with zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in files:
                zf.write(f, arcname=f.name)
    _update(jid, status="done", progress=100,
            detail="Listo. Guardado en tu carpeta.",
            filename=final.name, filepath=str(final))
    _notify("KNAVE · Listo", final.name)
    if not SERVER_MODE and load_settings(sid)["auto_open_folder"]:
        _open_in_manager(final, select=True)


# ----------------------------------------------------------------------------
# ffmpeg — junto al ejecutable (app empaquetada) o en el PATH
# ----------------------------------------------------------------------------
def _find_ffmpeg() -> str | None:
    exe_dir = Path(sys.executable).parent if getattr(sys, "frozen", False) else BUNDLE_DIR
    for cand in (exe_dir / "ffmpeg.exe", exe_dir / "ffmpeg",
                 BUNDLE_DIR / "ffmpeg.exe", BUNDLE_DIR / "ffmpeg"):
        if cand.exists():
            return str(cand)
    return shutil.which("ffmpeg")


FFMPEG = _find_ffmpeg()


# ----------------------------------------------------------------------------
# Abrir el explorador de archivos del sistema (nunca debe romper nada)
# ----------------------------------------------------------------------------
def _open_in_manager(path: Path, select: bool = False) -> None:
    try:
        if sys.platform == "win32":
            if select and path.is_file():
                subprocess.Popen(["explorer", "/select,", str(path)])
            else:
                os.startfile(str(path if path.is_dir() else path.parent))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", "-R", str(path)] if select else ["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path if path.is_dir() else path.parent)])
    except Exception:  # noqa: BLE001
        pass


# ----------------------------------------------------------------------------
# Descargas — yt-dlp (YouTube y cientos de sitios) y spotdl (Spotify)
# ----------------------------------------------------------------------------
def _is_spotify(url: str) -> bool:
    if url.startswith("spotify:"):
        return True
    host = (urlparse(url).hostname or "").lower()
    return host == "spotify.com" or host.endswith(".spotify.com")


def preview_url(url: str) -> dict:
    """Metadatos ligeros antes de descargar (mejor esfuerzo)."""
    if _is_spotify(url):
        return {"kind": "spotify", "title": "Enlace de Spotify",
                "detail": "Se descargará como audio con metadatos de Spotify.",
                "is_playlist": "/playlist/" in url or "/album/" in url}
    import yt_dlp
    opts = {"quiet": True, "no_warnings": True, "skip_download": True,
            "extract_flat": "in_playlist", "socket_timeout": 20,
            "playlist_items": "1-1"}
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(url, download=False)
    if info.get("_type") == "playlist" or "entries" in info:
        entries = list(info.get("entries") or [])
        first = entries[0] if entries else {}
        thumbs = info.get("thumbnails") or first.get("thumbnails") or []
        return {
            "kind": "playlist",
            "title": info.get("title") or "Lista de reproducción",
            "uploader": info.get("uploader") or info.get("channel") or "",
            "count": info.get("playlist_count") or len(entries),
            "thumbnail": (thumbs[-1].get("url") if thumbs else None),
            "is_playlist": True,
        }
    thumbs = info.get("thumbnails") or []
    return {
        "kind": "video",
        "title": info.get("title") or "Vídeo",
        "uploader": info.get("uploader") or info.get("channel") or "",
        "duration": info.get("duration"),
        "thumbnail": info.get("thumbnail") or (thumbs[-1].get("url") if thumbs else None),
        "is_playlist": False,
    }


def _run_ytdlp(jid: str, url: str, mode: str, quality: str, fmt: str,
               playlist: bool, subs: bool, embed: bool, outdir: Path,
               ck: Path | None = None) -> None:
    import yt_dlp

    def hook(d: dict) -> None:
        if _is_cancelled(jid):
            raise Cancelled()
        if d.get("status") == "downloading":
            total = d.get("total_bytes") or d.get("total_bytes_estimate") or 0
            done = d.get("downloaded_bytes") or 0
            pct = int(done * 100 / total) if total else 0
            speed = (d.get("_speed_str") or "").strip()
            info = d.get("info_dict") or {}
            idx, n = info.get("playlist_index"), info.get("n_entries")
            tag = f"[{idx}/{n}] " if idx and n else ""
            _update(jid, status="running", progress=min(pct, 99),
                    detail=f"{tag}Descargando… {min(pct, 99)}%"
                           + (f" · {speed}" if speed else ""))
        elif d.get("status") == "finished":
            _update(jid, progress=99, detail="Procesando con ffmpeg…")

    tmpl = ("%(playlist_index)02d - %(title).150B.%(ext)s" if playlist
            else "%(title).170B.%(ext)s")
    opts: dict = {
        "outtmpl": str(outdir / tmpl),
        "noplaylist": not playlist,
        "retries": 10, "fragment_retries": 10, "socket_timeout": 30,
        "windowsfilenames": True, "quiet": True, "no_warnings": True,
        "noprogress": True, "progress_hooks": [hook],
        "concurrent_fragment_downloads": 4,
        "postprocessors": [],
        "ignoreerrors": playlist,   # en listas, una entrada rota no tumba el resto
    }
    if not playlist:
        opts["playlist_items"] = "1"        # un enlace de lista pura → solo el 1.º
    if FFMPEG:
        opts["ffmpeg_location"] = FFMPEG

    if mode == "audio":
        codec = fmt if fmt in AUDIO_FORMATS else "mp3"
        q = quality if quality in AUDIO_QUALITIES else "192"
        opts["format"] = "bestaudio/best"
        opts["postprocessors"].append({
            "key": "FFmpegExtractAudio", "preferredcodec": codec,
            "preferredquality": q,
        })
        if embed:
            opts["postprocessors"].append({"key": "FFmpegMetadata",
                                           "add_metadata": True})
            if codec in {"mp3", "m4a", "flac", "opus"}:   # wav no admite carátula
                opts["writethumbnail"] = True
                opts["postprocessors"].append({"key": "EmbedThumbnail"})
    else:
        if quality in VIDEO_QUALITIES and quality != "best":
            opts["format"] = (f"bestvideo[height<={quality}]+bestaudio/"
                              f"best[height<={quality}]/best")
        else:
            opts["format"] = "bestvideo*+bestaudio/best"
        opts["merge_output_format"] = "mp4"
        opts["postprocessors"].append({"key": "FFmpegMetadata", "add_metadata": True})
        if subs:
            opts["writesubtitles"] = True
            opts["writeautomaticsub"] = True
            opts["subtitleslangs"] = ["es", "en"]
            opts["postprocessors"].append({"key": "FFmpegEmbedSubtitle"})

    if ck is not None and ck.exists():
        opts["cookiefile"] = str(ck)

    # Cascada de rescate: si una vía falla, se intenta la siguiente antes de
    # rendirse. 2ª vía: otros clientes de YouTube. 3ª: formato más permisivo.
    strategies: list[dict] = [
        {},
        {"extractor_args":
            {"youtube": {"player_client": ["android", "web_safari", "tv"]}}},
        {"format": ("bestaudio/best" if mode == "audio"
                    else "best[ext=mp4]/best[height<=1080]/best"),
         "postprocessors": (opts["postprocessors"][:1] if mode == "audio"
                            else [])},
    ]
    last_error: Exception | None = None
    for n, extra in enumerate(strategies, 1):
        if n > 1:
            _update(jid, progress=2,
                    detail=f"La vía {n - 1} falló; probando otra ({n}/{len(strategies)})…")
        try:
            with yt_dlp.YoutubeDL({**opts, **extra}) as ydl:
                ydl.download([url])
            return
        except Cancelled:
            raise
        except Exception as exc:  # noqa: BLE001 — a por la siguiente vía
            last_error = exc
    raise last_error if last_error else RuntimeError("Descarga imposible.")


def _run_spotdl(jid: str, url: str, quality: str, fmt: str, outdir: Path,
                ck: Path | None = None) -> None:
    """Spotify vía la API de spotdl (compatible con la app empaquetada):
    metadatos de Spotify + audio equivalente desde YouTube (sin tocar DRM)."""
    from spotdl import Spotdl
    from spotdl.utils.config import DEFAULT_CONFIG

    codec = fmt if fmt in {"mp3", "flac", "opus", "m4a", "wav", "ogg"} else "mp3"
    settings = {
        "output": str(outdir / "{artists} - {title}.{output-ext}"),
        "format": codec,
        "bitrate": f"{quality}k" if quality in AUDIO_QUALITIES else "192k",
        "simple_tui": True,
    }
    if FFMPEG:
        settings["ffmpeg"] = FFMPEG
    if ck is not None and ck.exists():
        settings["cookie_file"] = str(ck)

    _update(jid, status="running", detail="Buscando metadatos en Spotify…", progress=4)
    spot = Spotdl(client_id=DEFAULT_CONFIG["client_id"],
                  client_secret=DEFAULT_CONFIG["client_secret"],
                  downloader_settings=settings)
    songs = spot.search([url])
    if not songs:
        raise RuntimeError("Spotify no devolvió ninguna pista. "
                           "Comprueba que el enlace sea público y válido.")
    total, failed = len(songs), 0
    for i, song in enumerate(songs, 1):
        _check_cancel(jid)
        _update(jid, progress=5 + int(90 * (i - 1) / total),
                detail=f"Pista {i} de {total} · {song.display_name[:70]}")
        try:
            _, path = spot.download(song)
            if path is None:
                failed += 1
        except Exception:  # noqa: BLE001 — una pista rota no tumba el lote
            failed += 1
    if failed == total:
        raise RuntimeError("No se pudo descargar ninguna pista de ese enlace.")


def _download_worker(jid: str, url: str, mode: str, quality: str, fmt: str,
                     playlist: bool, subs: bool, embed: bool) -> None:
    outdir = _job_dir(jid)
    with JOBS_LOCK:
        sid = JOBS.get(jid, {}).get("sid", "local")
    ck = _cookies_path(sid)
    watchdog = threading.Timer(DOWNLOAD_TIMEOUT, lambda: _fail(
        jid, "Tiempo agotado: la descarga tardó demasiado y fue cancelada."))
    watchdog.daemon = True
    watchdog.start()
    try:
        _update(jid, status="running", detail="Iniciando descarga…", progress=1)
        if _is_spotify(url):
            _run_spotdl(jid, url, quality, fmt, outdir, ck)
        else:
            _run_ytdlp(jid, url, mode, quality, fmt, playlist, subs, embed,
                       outdir, ck)
        _check_cancel(jid)

        def _keep(p: Path) -> bool:
            if not p.is_file() or p.name.endswith((".part", ".ytdl", ".tmp")):
                return False
            if mode == "audio" and p.suffix.lower() in {".webp", ".jpg", ".jpeg", ".png"}:
                return False        # miniaturas sueltas tras incrustar la carátula
            return True

        files = sorted((p for p in outdir.rglob("*") if _keep(p)),
                       key=lambda p: p.stat().st_mtime)
        _publish(jid, files, "descarga_knave")
    except Cancelled:
        _update(jid, status="cancelled", detail="Cancelado.", progress=100)
    except Exception as exc:  # noqa: BLE001
        msg = re.sub(r"\x1b\[[0-9;]*m", "", str(exc) or exc.__class__.__name__)
        msg = msg.replace("ERROR: ", "").strip()[:400]
        _fail(jid, f"No se pudo descargar: {msg}")
    finally:
        watchdog.cancel()
        shutil.rmtree(outdir, ignore_errors=True)


# ----------------------------------------------------------------------------
# Conversiones
# ----------------------------------------------------------------------------
def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        if (p := shutil.which(name)):
            return p
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(c).exists():
            return c
    return None


def _run_cancellable(jid: str, cmd: list[str], timeout: int) -> str:
    """Proceso externo con timeout y cancelación por el usuario."""
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE,
                            stderr=subprocess.STDOUT, text=True)
    _update(jid, proc=proc)
    deadline = time.time() + timeout
    try:
        while proc.poll() is None:
            if _is_cancelled(jid):
                proc.kill()
                raise Cancelled()
            if time.time() > deadline:
                proc.kill()
                raise RuntimeError("La operación externa tardó demasiado.")
            time.sleep(0.2)
        out = proc.stdout.read() if proc.stdout else ""
    finally:
        _update(jid, proc=None)
    return out.strip()[-300:]


def _office_to_pdf(jid: str, src: Path, outdir: Path) -> Path:
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice no está instalado o no se encuentra. "
                           "Instálalo para convertir documentos a PDF.")
    profile = outdir / ".lo_profile"
    cmd = [soffice, "--headless", "--norestore", "--nolockcheck",
           f"-env:UserInstallation=file://{profile.as_posix()}",
           "--convert-to", "pdf", "--outdir", str(outdir), str(src)]
    _update(jid, progress=30, detail="Convirtiendo con LibreOffice…")
    last = ""
    for attempt in (1, 2):    # LibreOffice a veces falla en frío; un reintento
        last = _run_cancellable(jid, cmd, SOFFICE_TIMEOUT)
        produced = outdir / (src.stem + ".pdf")
        if produced.exists() and produced.stat().st_size > 0:
            return produced
        _update(jid, detail=f"Reintentando conversión… (intento {attempt + 1})")
        time.sleep(1.5)
    raise RuntimeError("LibreOffice no generó el PDF. "
                       + (f"Detalle: {last}" if last else ""))


def _pdf_to_docx(jid: str, src: Path, dst: Path) -> None:
    from pdf2docx import Converter
    _update(jid, progress=25, detail="Reconstruyendo el documento en Word…")
    cv = Converter(str(src))
    try:
        cv.convert(str(dst))
    finally:
        cv.close()
    if not dst.exists() or dst.stat().st_size == 0:
        raise RuntimeError("No se pudo reconstruir el PDF como documento Word.")


def _open_pdf(src: Path):
    import fitz
    doc = fitz.open(str(src))
    if doc.page_count == 0:
        doc.close()
        raise RuntimeError("El PDF no tiene páginas.")
    if doc.page_count > MAX_PDF_PAGES:
        n = doc.page_count
        doc.close()
        raise RuntimeError(f"El PDF tiene {n} páginas; el máximo es {MAX_PDF_PAGES}.")
    return doc


def _pdf_to_pptx(jid: str, src: Path, dst: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu
    doc = _open_pdf(src)
    try:
        pages = doc.page_count
        first = doc.load_page(0).rect
        prs = Presentation()
        prs.slide_width = Emu(int(first.width * 12700))     # 1 pt = 12700 EMU
        prs.slide_height = Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        for i in range(pages):
            _check_cancel(jid)
            pix = doc.load_page(i).get_pixmap(dpi=150)
            img = dst.parent / f"__p{i:04d}.png"
            pix.save(str(img))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(img), 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
            img.unlink(missing_ok=True)
            _update(jid, progress=10 + int(80 * (i + 1) / pages),
                    detail=f"Página {i + 1} de {pages}…")
        prs.save(str(dst))
    finally:
        doc.close()


def _pdf_to_images(jid: str, src: Path, outdir: Path, stem: str) -> list[Path]:
    doc = _open_pdf(src)
    made: list[Path] = []
    try:
        pages = doc.page_count
        for i in range(pages):
            _check_cancel(jid)
            pix = doc.load_page(i).get_pixmap(dpi=200)
            img = outdir / f"{stem}_pagina_{i + 1:03d}.png"
            pix.save(str(img))
            made.append(img)
            _update(jid, progress=10 + int(85 * (i + 1) / pages),
                    detail=f"Página {i + 1} de {pages}…")
    finally:
        doc.close()
    return made


def _pdf_split(jid: str, src: Path, outdir: Path, stem: str) -> list[Path]:
    import fitz
    doc = _open_pdf(src)
    made: list[Path] = []
    try:
        pages = doc.page_count
        for i in range(pages):
            _check_cancel(jid)
            one = fitz.open()
            one.insert_pdf(doc, from_page=i, to_page=i)
            out = outdir / f"{stem}_pagina_{i + 1:03d}.pdf"
            one.save(str(out))
            one.close()
            made.append(out)
            _update(jid, progress=10 + int(85 * (i + 1) / pages),
                    detail=f"Página {i + 1} de {pages}…")
    finally:
        doc.close()
    return made


def _pdf_compress(jid: str, src: Path, dst: Path) -> str:
    """Compresión sin pérdida de texto (deflate + limpieza de basura)."""
    _update(jid, progress=40, detail="Optimizando el PDF…")
    doc = _open_pdf(src)
    try:
        doc.save(str(dst), garbage=4, deflate=True, deflate_images=True,
                 deflate_fonts=True, clean=True)
    finally:
        doc.close()
    before, after = src.stat().st_size, dst.stat().st_size
    if after >= before:
        shutil.copyfile(src, dst)
        return "El PDF ya estaba optimizado; se guardó una copia."
    saved = 100 * (before - after) / before
    return f"Reducido un {saved:.0f}% ({before // 1024} KB → {after // 1024} KB)."


def _img_to_pdf(jid: str, srcs: list[Path], dst: Path) -> None:
    import img2pdf
    from PIL import Image
    _update(jid, progress=30, detail=f"Uniendo {len(srcs)} imagen(es) en un PDF…")
    prepared: list[str] = []
    tmp: list[Path] = []
    for i, p in enumerate(srcs):
        _check_cancel(jid)
        # img2pdf no acepta alfa ni paletas: se normaliza a RGB con Pillow
        with Image.open(p) as im:
            if im.mode in ("RGBA", "LA", "P"):
                bg = Image.new("RGB", im.size, (255, 255, 255))
                im = im.convert("RGBA")
                bg.paste(im, mask=im.split()[-1])
                fixed = dst.parent / f"__img{i:03d}.jpg"
                bg.save(fixed, "JPEG", quality=92)
                prepared.append(str(fixed))
                tmp.append(fixed)
            else:
                prepared.append(str(p))
    with dst.open("wb") as fh:
        fh.write(img2pdf.convert(prepared))
    for t in tmp:
        t.unlink(missing_ok=True)


def _pdf_merge(jid: str, srcs: list[Path], dst: Path) -> None:
    import fitz
    merged = fitz.open()
    try:
        total = len(srcs)
        for i, p in enumerate(srcs, 1):
            _check_cancel(jid)
            with fitz.open(str(p)) as d:
                merged.insert_pdf(d)
            _update(jid, progress=10 + int(80 * i / total),
                    detail=f"Añadiendo {i} de {total}…")
        merged.save(str(dst), garbage=3, deflate=True)
    finally:
        merged.close()


def _convert_worker(jid: str, op: str, srcs: list[Path], first_stem: str) -> None:
    workdir = srcs[0].parent
    try:
        _update(jid, status="running", progress=8, detail="Analizando…")
        note = ""
        if op == "office2pdf":
            lo_out = _office_to_pdf(jid, srcs[0], workdir)
            final = workdir / f"{first_stem}.pdf"
            if lo_out != final:
                lo_out.replace(final)
            produced = [final]
        elif op == "pdf2docx":
            out = workdir / f"{first_stem}.docx"
            _pdf_to_docx(jid, srcs[0], out); produced = [out]
        elif op == "pdf2pptx":
            out = workdir / f"{first_stem}.pptx"
            _pdf_to_pptx(jid, srcs[0], out); produced = [out]
        elif op == "pdf2img":
            produced = _pdf_to_images(jid, srcs[0], workdir, first_stem)
        elif op == "pdfsplit":
            produced = _pdf_split(jid, srcs[0], workdir, first_stem)
        elif op == "pdfcompress":
            out = workdir / f"{first_stem}_comprimido.pdf"
            note = _pdf_compress(jid, srcs[0], out); produced = [out]
        elif op == "img2pdf":
            out = workdir / f"{first_stem}.pdf"
            _img_to_pdf(jid, srcs, out); produced = [out]
        elif op == "pdfmerge":
            out = workdir / "combinado.pdf"
            _pdf_merge(jid, srcs, out); produced = [out]
        else:
            raise RuntimeError(f"Operación desconocida: {op}.")

        zip_stem = {"pdf2img": f"{first_stem}_imagenes",
                    "pdfsplit": f"{first_stem}_paginas"}.get(op, first_stem)
        _publish(jid, produced, zip_stem)
        if note:
            with JOBS_LOCK:
                if JOBS.get(jid, {}).get("status") == "done":
                    JOBS[jid]["detail"] = note
    except Cancelled:
        _update(jid, status="cancelled", detail="Cancelado.", progress=100)
    except Exception as exc:  # noqa: BLE001
        _fail(jid, f"No se pudo convertir: {str(exc)[:400]}")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


# ----------------------------------------------------------------------------
# API — descargas
# ----------------------------------------------------------------------------
@app.post("/api/preview")
async def api_preview(payload: dict):
    url = str(payload.get("url", "")).strip()
    if not url:
        raise HTTPException(400, "Pega un enlace.")
    try:
        return preview_url(url)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f"No se pudo leer el enlace: {str(exc)[:200]}")


@app.post("/api/jobs/download")
async def create_download(request: Request, payload: dict):
    raw = str(payload.get("url", "")).strip()
    urls = [u.strip() for u in raw.splitlines() if u.strip()]
    if not urls:
        raise HTTPException(400, "Pega al menos un enlace.")
    if len(urls) > 20:
        raise HTTPException(400, "Máximo 20 enlaces por tanda.")

    sid = request.state.sid
    s = load_settings(sid)
    mode = payload.get("mode", "video")
    quality = str(payload.get("quality", "best"))
    fmt = str(payload.get("audio_format", s["audio_format"]))
    playlist = bool(payload.get("playlist", False))
    subs = bool(payload.get("subs", s["download_subs"]))
    embed = bool(payload.get("embed", s["embed_thumbnail"]))
    if mode not in {"audio", "video"}:
        raise HTTPException(400, "Modo inválido: usa 'audio' o 'video'.")

    job_ids: list[str] = []
    for url in urls:
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"} and not url.startswith("spotify:"):
            raise HTTPException(400, f"Enlace no válido: {url[:60]}")
        m = "audio" if _is_spotify(url) else mode
        label = (("Audio" if m == "audio" else "Vídeo") + " · "
                 + (parsed.hostname or "spotify"))
        jid = _new_job("download", label, sid)
        POOL.submit(_download_worker, jid, url, m, quality, fmt,
                    playlist, subs, embed)
        job_ids.append(jid)
    return {"job_ids": job_ids}


# ----------------------------------------------------------------------------
# API — conversiones
# ----------------------------------------------------------------------------
@app.get("/api/operations")
async def api_operations():
    return {k: {"label": v["label"], "min": v["min"], "max": v["max"],
                "exts": sorted(v["exts"])} for k, v in OPERATIONS.items()}


@app.post("/api/jobs/convert")
async def create_convert(request: Request, op: str = Form(...),
                         files: list[UploadFile] = File(...)):
    spec = OPERATIONS.get(op)
    if not spec:
        raise HTTPException(400, f"Operación desconocida: {op}.")
    if not (spec["min"] <= len(files) <= spec["max"]):
        raise HTTPException(400, f"«{spec['label']}» admite entre {spec['min']} y "
                                 f"{spec['max']} archivo(s); recibí {len(files)}.")

    jid = _new_job("convert", f"{spec['label']} · {len(files)} archivo(s)",
                   request.state.sid)
    workdir = _job_dir(jid)
    saved: list[Path] = []
    limit = MAX_UPLOAD_MB * 1024 * 1024
    try:
        for idx, up in enumerate(files):
            name = _sanitize_name(up.filename or f"archivo_{idx}")
            ext = Path(name).suffix.lower().lstrip(".")
            if ext not in spec["exts"]:
                raise HTTPException(400, f"«{spec['label']}» no acepta .{ext or '?'}.")
            dest = workdir / f"{idx:03d}_{name}"     # el prefijo conserva el orden
            size = 0
            with dest.open("wb") as fh:
                while chunk := await up.read(1024 * 1024):
                    size += len(chunk)
                    if size > limit:
                        raise HTTPException(
                            413, f"«{name}» supera el límite de {MAX_UPLOAD_MB} MB.")
                    fh.write(chunk)
            if size == 0:
                raise HTTPException(400, f"«{name}» llegó vacío.")
            saved.append(dest)
    except HTTPException:
        shutil.rmtree(workdir, ignore_errors=True)
        with JOBS_LOCK:
            JOBS.pop(jid, None)
        raise

    first_stem = _sanitize_name(Path(re.sub(r"^\d{3}_", "", saved[0].name)).stem)
    POOL.submit(_convert_worker, jid, op, saved, first_stem)
    return {"job_id": jid}


# ----------------------------------------------------------------------------
# API — trabajos
# ----------------------------------------------------------------------------
@app.get("/api/jobs/{jid}")
async def job_status(jid: str, request: Request):
    with JOBS_LOCK:
        job = JOBS.get(jid)
        if job is None or job.get("sid") != request.state.sid:
            raise HTTPException(404, "Ese trabajo no existe o ya expiró.")
        return {k: v for k, v in job.items()
                if k not in ("filepath", "proc", "sid")}


@app.post("/api/jobs/{jid}/cancel")
async def cancel_job(jid: str, request: Request):
    with JOBS_LOCK:
        job = JOBS.get(jid)
        if job is None or job.get("sid") != request.state.sid:
            raise HTTPException(404, "Ese trabajo no existe.")
        if job["status"] in ("done", "error", "cancelled"):
            return {"ok": True, "status": job["status"]}
        job["cancelled"] = True
        proc = job.get("proc")
    if proc is not None and proc.poll() is None:
        proc.kill()
    return {"ok": True, "status": "cancelling"}


@app.get("/api/jobs/{jid}/file")
async def job_file(jid: str, request: Request):
    with JOBS_LOCK:
        job = JOBS.get(jid)
    if (job is None or job.get("sid") != request.state.sid
            or job["status"] != "done" or not job.get("filepath")):
        raise HTTPException(404, "El archivo no está listo.")
    path = Path(job["filepath"])
    if not path.exists():
        raise HTTPException(410, "El archivo ya no está en su ubicación.")
    media = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return FileResponse(path, filename=job["filename"], media_type=media)


# ----------------------------------------------------------------------------
# API — biblioteca (los archivos permanentes de la carpeta de salida)
# ----------------------------------------------------------------------------
@app.get("/api/library")
async def api_library(request: Request):
    out = output_dir(request.state.sid)
    items = []
    for p in out.iterdir():
        if p.is_file() and not p.name.startswith("."):
            st = p.stat()
            items.append({"name": p.name, "size": st.st_size, "mtime": st.st_mtime,
                          "ext": p.suffix.lower().lstrip(".")})
    items.sort(key=lambda x: x["mtime"], reverse=True)
    return {"dir": str(out), "items": items[:200]}


def _safe_in_output(name: str, sid: str) -> Path:
    if not name or name in (".", ".."):
        raise HTTPException(400, "Nombre no válido.")
    out = output_dir(sid).resolve()
    target = (out / name).resolve()
    if out not in target.parents:
        raise HTTPException(400, "Ruta no permitida.")
    if not target.exists() or not target.is_file():
        raise HTTPException(404, "El archivo no existe.")
    return target


@app.get("/api/library/file")
async def library_file(name: str, request: Request):
    path = _safe_in_output(name, request.state.sid)
    media = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return FileResponse(path, filename=path.name, media_type=media)


@app.post("/api/library/reveal")
async def library_reveal(request: Request, payload: dict):
    if SERVER_MODE:
        raise HTTPException(400, "No disponible en la versión web.")
    path = _safe_in_output(str(payload.get("name", "")), request.state.sid)
    _open_in_manager(path, select=True)
    return {"ok": True}


@app.post("/api/library/delete")
async def library_delete(request: Request, payload: dict):
    path = _safe_in_output(str(payload.get("name", "")), request.state.sid)
    path.unlink(missing_ok=True)
    return {"ok": True}


@app.post("/api/open-folder")
async def open_folder():
    if SERVER_MODE:
        raise HTTPException(400, "No disponible en la versión web.")
    _open_in_manager(output_dir(), select=False)
    return {"ok": True}


# ----------------------------------------------------------------------------
# API — ajustes y estado
# ----------------------------------------------------------------------------
@app.get("/api/settings")
async def get_settings(request: Request):
    s = load_settings(request.state.sid)
    s["has_cookies"] = _cookies_path(request.state.sid).exists()
    return s


@app.post("/api/settings")
async def update_settings(request: Request, payload: dict):
    if SERVER_MODE:
        payload.pop("output_dir", None)   # en la web la carpeta es fija por sesión
    if "output_dir" in payload:
        d = str(payload["output_dir"]).strip()
        if d:
            try:
                Path(d).expanduser().mkdir(parents=True, exist_ok=True)
            except Exception:  # noqa: BLE001
                raise HTTPException(400, "No se pudo crear esa carpeta de salida.")
    return save_settings(payload, request.state.sid)


@app.post("/api/cookies")
async def upload_cookies(request: Request, file: UploadFile = File(...)):
    data = await file.read()
    if not data or len(data) > 1024 * 1024:
        raise HTTPException(400, "El archivo de cookies llegó vacío o es demasiado grande.")
    if "\t" not in data.decode("utf-8", "ignore"):
        raise HTTPException(400, "Eso no parece un cookies.txt en formato Netscape.")
    _cookies_path(request.state.sid).write_bytes(data)
    return {"ok": True}


@app.post("/api/cookies/delete")
async def remove_cookies(request: Request):
    _cookies_path(request.state.sid).unlink(missing_ok=True)
    return {"ok": True}


@app.get("/api/health")
async def health(request: Request):
    return {"ok": True, "server": SERVER_MODE, "ffmpeg": FFMPEG,
            "libreoffice": _find_soffice() is not None,
            "output_dir": str(output_dir(request.state.sid))}


@app.get("/healthz")
async def healthz():
    return {"ok": True}


# ----------------------------------------------------------------------------
# Limpieza del borrador + frontend estático
# ----------------------------------------------------------------------------
def _janitor() -> None:
    while True:
        time.sleep(900)
        cutoff = time.time() - SCRATCH_TTL_SECONDS
        with JOBS_LOCK:
            expired = [j for j, job in JOBS.items() if job["created"] < cutoff
                       and job["status"] in ("done", "error", "cancelled")]
            for j in expired:
                JOBS.pop(j, None)
        for d in WORKSPACE.iterdir():
            try:
                if d.is_dir() and d.stat().st_mtime < cutoff:
                    shutil.rmtree(d, ignore_errors=True)
            except OSError:
                pass
        if SERVER_MODE:
            ttl = time.time() - SESSION_TTL_HOURS * 3600
            for sess in list(SESSIONS_DIR.iterdir()):
                lib = sess / "library"
                try:
                    if lib.is_dir():
                        for f in lib.iterdir():
                            if f.is_file() and f.stat().st_mtime < ttl:
                                f.unlink(missing_ok=True)
                    if sess.stat().st_mtime < ttl and (
                            not lib.is_dir() or not any(lib.iterdir())):
                        shutil.rmtree(sess, ignore_errors=True)
                except OSError:
                    pass


threading.Thread(target=_janitor, daemon=True).start()

app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", "8666"))
    host = "0.0.0.0" if SERVER_MODE else "127.0.0.1"
    print(f"\n  KNAVE encendida →  http://127.0.0.1:{port}"
          + ("  (modo servidor)" if SERVER_MODE else "") + "\n")
    uvicorn.run(app, host=host, port=port, log_level="warning")
